import base64
import contextlib
import csv
import hashlib
import hmac
import http.client
import io
import ipaddress
import json
import os
import re
import secrets
import socket
import threading
import time
import urllib.parse
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime
from decimal import Decimal
from typing import Literal

import jwt
import pillow_heif
import pypdfium2
from cryptography.fernet import Fernet, InvalidToken
from PIL import Image, ImageOps
from fastapi import BackgroundTasks, Depends, FastAPI, Header, HTTPException, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response
from psycopg.types.json import Jsonb
from pydantic import BaseModel, Field

import providers
import store
from schema import Document
from validate import apply_date_order, suggest, validate

providers.load_env()
pillow_heif.register_heif_opener()   # lets Pillow open iPhone HEIC photos
MODEL = os.environ.get('EXTRACT_MODEL', 'gemini-3.1-flash-lite')
MAX_BYTES = 10 * 1024 * 1024
MAX_FILES = 20
MAX_PDF_PAGES = 20   # also caps model cost: the whole PDF goes to the model
PDF_LOCK = threading.Lock()   # PDFium is not thread-safe; endpoints run in a thread pool
DateOrderValue = Literal['MDY', 'DMY']
DAILY_UPLOAD_LIMIT = int(os.environ.get('DAILY_UPLOAD_LIMIT', 10))   # model reads (uploads + retries) per day: protects the quota
SECOND_MODEL = os.environ.get('SECOND_MODEL', 'gemini-3.5-flash')
SECOND_READ_DAILY_LIMIT = int(os.environ.get('SECOND_READ_DAILY_LIMIT', 15))   # whole app, any 24 hours: the free quota is shared
SECOND_READ_PER_USER = int(os.environ.get('SECOND_READ_PER_USER', 5))   # per user, any 24 hours, on top of the app-wide cap
SECOND_READ_FLAGGED_RESERVE = 5   # passed receipts use only leftover quota: this many stay for flagged ones
REAL_CHECK_EXCLUDED = {'date_ambiguous', 'duplicate'}   # a second reading cannot settle these


def reads_today(con, uid):
    return con.execute("SELECT count(*) AS n FROM reads WHERE user_id = %s AND at > now() - interval '1 day'", (uid,)).fetchone()['n']

@contextlib.asynccontextmanager
async def lifespan(_):
    threading.Thread(target=resume_stuck, daemon=True).start()
    threading.Thread(target=deliver_events, daemon=True).start()  # also sends events left over from before a restart
    yield


app = FastAPI(title='Crosscheck API', lifespan=lifespan)
# the web app calls the API from the browser; only its own origin(s) may (comma-separated FRONTEND_ORIGIN)
app.add_middleware(CORSMiddleware, allow_origins=[o.strip().rstrip('/') for o in os.environ.get('FRONTEND_ORIGIN', 'http://localhost:3000').split(',')],
                   allow_methods=['*'], allow_headers=['Authorization', 'Content-Type'], expose_headers=['X-Skipped', 'Content-Disposition'])  # export file names
_jwks = None


def signing_key(token):
    global _jwks
    if _jwks is None:
        _jwks = jwt.PyJWKClient(f'{os.environ["SUPABASE_URL"].rstrip("/")}/auth/v1/.well-known/jwks.json')
    return _jwks.get_signing_key_from_jwt(token).key


def current_user(authorization: str | None = Header(None)) -> str:
    if not authorization or not authorization.startswith('Bearer '):
        raise HTTPException(401, 'Please sign in.')
    token = authorization.removeprefix('Bearer ')
    try:
        claims = jwt.decode(token, signing_key(token), algorithms=['ES256', 'RS256'], audience='authenticated', options={'require': ['exp', 'sub']}, leeway=30,  # the PC's clock may run a little behind Supabase's
                            issuer=f'{os.environ["SUPABASE_URL"].rstrip("/")}/auth/v1')
    except jwt.PyJWTError:
        raise HTTPException(401, 'Please sign in again.')
    return claims['sub']


COLUMNS = 'id, user_id, file_name, mime, status, document, checks, error, vendor, currency, issue_date, total, model, ' \
          'prompt_version, tokens_in, tokens_out, created_at, updated_at, second_reading'


def row_by_id(con, doc_id, with_file=False):
    # no owner check: background jobs only; endpoints use get_row
    return con.execute(f'SELECT {COLUMNS}{", file" if with_file else ""} FROM documents WHERE id = %s', (doc_id,)).fetchone()


def get_row(con, doc_id, uid, with_file=False):
    # 404 (not 403) for other users' documents, so their ids reveal nothing
    r = row_by_id(con, doc_id, with_file)
    if r is None or str(r['user_id']) != uid:
        raise HTTPException(404, 'Document not found.')
    return r


def run_checks(con, uid, doc: Document, order, doc_id=None):
    checks = validate(doc, date_order=order)
    dup = store.duplicate_of(con, uid, doc, before_id=doc_id)
    if dup:
        checks.append({'check': 'duplicate', 'fields': ['doc_number'], 'duplicate_of': dup['id'],
                       'message': f'Same vendor, number and total as {dup["doc_number"]}. It was uploaded earlier.'})
    if doc_id is not None:
        row = con.execute('SELECT second_reading FROM documents WHERE id = %s AND user_id = %s', (doc_id, uid)).fetchone()
        if row and row['second_reading']:
            second = apply_date_order(Document(**row['second_reading']), order)
            changes = meaningful_changes(doc, second)
            if changes:
                checks.append({'check': 'second_reading', 'fields': sorted({c['field'].split('[')[0] for c in changes}),
                               'message': 'A second AI reading differs. Compare with the photo.'})
    return checks


def save(con, doc_id, doc: Document, status, uid, extra=None, date_order=None, only_if_processing=False):
    # only_if_processing: the background reader never overwrites corrections saved meanwhile
    order = date_order or store.date_order(con, uid, doc.vendor)
    doc = apply_date_order(doc, order)
    checks = run_checks(con, uid, doc, order, doc_id)
    if status is None:
        status = 'needs_review' if checks else 'passed'
    fields = {'document': Jsonb(doc.model_dump(mode='json')), 'checks': Jsonb(checks), 'status': status,
              'error': None, 'vendor': doc.vendor, 'currency': doc.currency, 'issue_date': doc.issue_date,
              'total': doc.total, **(extra or {})}
    if only_if_processing:  # the AI's reading: also kept apart, as the user first sees it, to measure corrections (corrections.py)
        fields['extracted'] = fields['document']
    only = " AND status = 'processing'" if only_if_processing else ''
    con.execute(f'UPDATE documents SET {", ".join(k + " = %s" for k in fields)}, updated_at = now() WHERE id = %s{only}',
                (*fields.values(), doc_id))


def webhook_users():  # comma-separated, so more than one of your own accounts can send events
    return {u.strip().lower() for u in os.environ.get('WEBHOOK_USER_ID', '').split(',') if u.strip()}


def fernet():
    # users' webhook addresses and secrets are stored encrypted with WEBHOOK_KEY (any long random text)
    key = os.environ.get('WEBHOOK_KEY')
    return key and Fernet(base64.urlsafe_b64encode(hashlib.sha256(key.encode()).digest()))


def webhook_for(con, uid):
    # (url, secret, users_own): the user's own address first, else the server's n8n for the accounts in WEBHOOK_USER_ID.
    # Read at send time, so waiting events go to the current address with the current secret.
    r = con.execute('SELECT url, secret FROM webhooks WHERE user_id = %s', (uid,)).fetchone()
    if r and fernet():
        try:
            return fernet().decrypt(r['url'].encode()).decode(), fernet().decrypt(r['secret'].encode()).decode(), True
        except InvalidToken:
            pass  # WEBHOOK_KEY changed: the user has to set the address again
    if os.environ.get('WEBHOOK_URL') and os.environ.get('WEBHOOK_SECRET') and str(uid) in webhook_users():
        return os.environ['WEBHOOK_URL'], os.environ['WEBHOOK_SECRET'], False
    return None  # a missing WEBHOOK_SECRET shows in the log as 'webhooks are not set up on the server'


def configured(con, uid):
    # the user set an address, or is one of the server's n8n accounts, even if webhook_for can't build the target now
    own = con.execute('SELECT 1 FROM webhooks WHERE user_id = %s', (uid,)).fetchone()
    return bool(own) or bool(os.environ.get('WEBHOOK_URL')) and str(uid) in webhook_users()


def public_ip(url):
    # Users' addresses must be https on the public internet, so the server can't be pointed at its own network.
    # Returns the IP to connect to (None: refused). Sending connects to exactly that IP, so the name can't
    # be pointed somewhere else between the check and the connection.
    try:
        p = urllib.parse.urlsplit(url)
        if p.scheme != 'https' or not p.hostname or p.username is not None:
            return None
        ips = [a[4][0] for a in socket.getaddrinfo(p.hostname, p.port or 443, proto=socket.IPPROTO_TCP)]
    except (OSError, UnicodeError, ValueError):
        return None
    return ips[0] if ips and all(ipaddress.ip_address(ip).is_global for ip in ips) else None


def public_url(url):
    return public_ip(url) is not None


def queue_event(con, uid, payload):
    # Written in the caller's transaction, so a saved change always has its event, even if the server stops right after.
    # The document row is locked by that change, so one document's events are queued in order.
    try:
        with con.transaction():  # savepoint: a failed insert never undoes the document change
            if configured(con, uid):  # kept even while the server's settings are broken, sent once they are fixed
                con.execute('INSERT INTO webhook_events (user_id, payload) VALUES (%s, %s)', (uid, Jsonb(payload)))
    except Exception as e:
        print(f'webhook event for document {payload["id"]} not queued: {e!r}')


def queue_document_event(con, doc_id):
    r = row_by_id(con, doc_id)
    if r:
        queue_event(con, r['user_id'], {'event': f'document.{r["status"]}', 'id': r['id'], 'file_name': r['file_name'],
                                        'status': r['status'], 'document': r['document'], 'checks': r['checks'], 'error': r['error']})


def signed(payload, secret):
    body = json.dumps(payload).encode()
    return body, {'Content-Type': 'application/json', 'User-Agent': 'receipt-extractor/0.1',
                  'X-Signature': 'sha256=' + hmac.new(secret.encode(), body, hashlib.sha256).hexdigest()}


def open_webhook(url, payload, secret, users_own, timeout):
    # http.client follows no redirects, so a redirect can't lead past the check
    p = urllib.parse.urlsplit(url)
    ip = public_ip(url) if users_own else None  # checked again at send time: the name may point elsewhere by now
    if users_own and not ip:
        raise ValueError('not a public https address')
    conn = (http.client.HTTPSConnection if p.scheme == 'https' else http.client.HTTPConnection)(p.hostname, p.port, timeout=timeout)
    if ip:  # certificate and Host header still use the name
        conn._create_connection = lambda address, *args: socket.create_connection((ip, address[1]), *args)

    def cut():  # `timeout` limits each wait; this limits the whole send, even if the answer comes one byte at a time
        with contextlib.suppress(OSError, AttributeError):
            conn.sock.shutdown(socket.SHUT_RDWR)

    deadline = threading.Timer(timeout, cut)
    deadline.start()
    try:
        conn.request('POST', urllib.parse.urlunsplit(('', '', p.path or '/', p.query, '')), *signed(payload, secret))
        status = conn.getresponse().status
    finally:
        deadline.cancel()
        conn.close()
    if status >= 300:
        raise RuntimeError(f'HTTP {status}')
    return status


RETRY_SECONDS = (60, 300, 1800, 7200)  # after the first try; then the event is dropped and the next one goes
WAKE = threading.Event()  # set after a commit that queued events
SENDING = set()  # users with an event in flight: one at a time per user keeps their events in order
SENDING_LOCK = threading.Lock()
SENDERS = ThreadPoolExecutor(8)  # a slow address holds one sender, never the others
SECOND_WORKER = ThreadPoolExecutor(1)  # second readings run one at a time, so a batch upload's first readings never wait on them


def deliver_events():
    # Assumes one API process. With several, claim rows with FOR UPDATE SKIP LOCKED.
    while True:
        WAKE.clear()
        try:
            wait = dispatch()
        except Exception as e:
            print(f'webhook dispatch failed: {e!r}')
            wait = 60
        WAKE.wait(wait)


def dispatch():
    # starts each user's oldest event when due; returns seconds until the next retry (60 at most, in case a wake was missed)
    with store.conn() as con:
        rows = con.execute('SELECT DISTINCT ON (user_id) id, user_id, '
                           'extract(epoch FROM next_at - now()) AS wait FROM webhook_events ORDER BY user_id, id').fetchall()
    waits = [60]
    for r in rows:
        if r['wait'] > 0:
            waits.append(float(r['wait']))
            continue
        with SENDING_LOCK:
            if r['user_id'] in SENDING:
                continue
            SENDING.add(r['user_id'])
        SENDERS.submit(deliver, r['id'], r['user_id'])
    return min(waits)


def deliver(event_id, uid):
    try:
        with store.conn() as con:
            # read again: a sender that just finished may have sent or rescheduled it after dispatch() looked
            r = con.execute('SELECT id, user_id, payload, attempts FROM webhook_events WHERE id = %s AND next_at <= now()',
                            (event_id,)).fetchone()
            target = r and webhook_for(con, uid)
            error = None
            if r and not target and configured(con, uid):
                error = 'webhooks are not set up on the server'  # e.g. WEBHOOK_KEY missing after a deploy: retried like a failed send, so it waits about 2.5 h
                print(f'webhook event {event_id} waits: {error}')
        if r is None:
            return
        if target:
            try:
                # 300 s for the server's n8n: asleep on Render's free plan it took 113 s to wake (measured 2026-10-02).
                # Users' own addresses get 30 s, so a few dead ones can't hold every sender; a sleeping one is awake by the retry.
                open_webhook(target[0], {**r['payload'], 'event_id': r['id']}, target[1], target[2], 30 if target[2] else 300)
            except Exception as e:
                error = str(e)[:200]
                print(f'webhook for document {r["payload"]["id"]} failed (attempt {r["attempts"] + 1}): {error}')
        with store.conn() as con:
            if error and r['attempts'] < len(RETRY_SECONDS):
                con.execute('UPDATE webhook_events SET attempts = attempts + 1, next_at = now() + make_interval(secs => %s) '
                            'WHERE id = %s', (RETRY_SECONDS[r['attempts']], r['id']))
            else:  # sent, given up, or nobody listens any more
                con.execute('DELETE FROM webhook_events WHERE id = %s', (r['id'],))
            if target and target[2]:  # shown on the Account page
                con.execute('UPDATE webhooks SET last_at = now(), last_error = %s WHERE user_id = %s', (error, r['user_id']))
    except Exception as e:
        print(f'webhook event {event_id} delivery failed: {e!r}')
    finally:
        with SENDING_LOCK:
            SENDING.discard(uid)
        WAKE.set()


def public_error(e):
    print(f'extraction failed: {e!r}')
    text = str(e)
    if 'daily quota used up' in text:
        return "Today's AI limit is used up. Try again tomorrow."
    if 'HTTP 503' in text or 'HTTP 429' in text:
        return 'The AI service is busy. Try again in a few minutes.'
    if 'no answer from the model service' in text:
        return 'The AI service did not answer. Try again.'
    if isinstance(e, ValueError):
        return 'The AI reply could not be read. Try again.'
    if re.search(r'HTTP (400|413|415)', text) and 'API key' not in text:  # the file itself; key errors stay generic
        return 'The AI could not read this file. Try a clearer photo or another file.'
    return 'Reading failed. Try again.'


def process(doc_id):
    try:
        with store.conn() as con:
            row = row_by_id(con, doc_id, with_file=True)
        if row is None:
            return  # deleted before extraction started
        text, tin, tout, _ = providers.call(MODEL, bytes(row['file']))
        doc = providers.parse(text)
        extra = {'model': MODEL, 'prompt_version': hashlib.sha256(providers.PROMPT.encode()).hexdigest()[:8],
                 'tokens_in': tin, 'tokens_out': tout}
        with store.conn() as con:
            save(con, doc_id, doc, None, str(row['user_id']), extra, only_if_processing=True)
            queue_document_event(con, doc_id)  # passed or needs_review (nothing if deleted meanwhile)
    except Exception as e:
        with store.conn() as con:
            failed = con.execute("UPDATE documents SET status = 'failed', error = %s, updated_at = now() "
                                 "WHERE id = %s AND status = 'processing'", (public_error(e), doc_id)).rowcount
            if failed:
                queue_document_event(con, doc_id)  # document.failed: an alert, the file could not be read
    WAKE.set()
    SECOND_WORKER.submit(second_read, doc_id)   # queued on its own single worker, so it never delays the next file's first reading


def second_read(doc_id):
    # a flagged document is read again, blind, by another model; a passed one too, while quota is left.
    # Its answer is offered in the review screen; on a passed document a meaningful disagreement sends it to review.
    # A document gets at most one second reading: a retry keeps it, since it read the same file blind and is still valid.
    try:
        with store.conn() as con:
            row = con.execute('SELECT status, checks, second_read_at, user_id, vendor, document FROM documents WHERE id = %s',
                              (doc_id,)).fetchone()
            if not row or row['second_read_at'] is not None:
                return
            passed = row['status'] == 'passed'
            flagged = row['status'] == 'needs_review' and {c['check'] for c in row['checks'] or []} - REAL_CHECK_EXCLUDED
            if not passed and not flagged:
                return
            uid = str(row['user_id'])
            per_user = con.execute("SELECT count(*) AS n FROM documents WHERE user_id = %s AND second_read_at > now() - interval '24 hours'",
                                   (uid,)).fetchone()['n']
            if per_user >= SECOND_READ_PER_USER:
                return
            used = con.execute("SELECT count(*) AS n FROM documents WHERE second_read_at > now() - interval '24 hours'").fetchone()['n']
            limit = SECOND_READ_DAILY_LIMIT - SECOND_READ_FLAGGED_RESERVE if passed else SECOND_READ_DAILY_LIMIT
            if used >= limit:  # exact: only one second reading runs at a time
                return
            file = con.execute('UPDATE documents SET second_read_at = now() WHERE id = %s RETURNING file', (doc_id,)).fetchone()['file']
        second = providers.parse(providers.call(SECOND_MODEL, bytes(file))[0])
        changed = False
        with store.conn() as con:
            con.execute('UPDATE documents SET second_reading = %s WHERE id = %s',
                        (Jsonb(second.model_dump(mode='json')), doc_id))
            if passed:
                order = store.date_order(con, uid, row['vendor'])
                doc = apply_date_order(Document(**row['document']), order)
                if meaningful_changes(doc, apply_date_order(second, order)):
                    checks = run_checks(con, uid, doc, order, doc_id)
                    changed = con.execute("UPDATE documents SET status = 'needs_review', checks = %s, updated_at = now() "
                                          "WHERE id = %s AND status = 'passed'", (Jsonb(checks), doc_id)).rowcount
                    if changed:
                        queue_document_event(con, doc_id)
        if changed:
            WAKE.set()
    except Exception as e:  # quota used up, busy, bad reply: the document simply goes to normal review
        print(f'second reading of document {doc_id} failed: {e!r}')


def resume_stuck():
    # documents still 'processing' at startup lost their job when the server stopped
    # Assumes one API process. With several, claim rows first (UPDATE ... RETURNING).
    try:
        with store.conn() as con:
            ids = [r['id'] for r in con.execute("SELECT id FROM documents WHERE status = 'processing' ORDER BY id")]
    except Exception as e:
        print(f'resuming stuck documents failed: {e!r}')
        return []
    for doc_id in ids:
        process(doc_id)   # never raises; failures are stored on the document
    return ids


@app.middleware('http')
async def limit_upload_size(request, call_next):
    # Relies on Content-Length: chunked uploads skip this check and spool to disk.
    if int(request.headers.get('content-length') or 0) > MAX_FILES * MAX_BYTES + 1024 * 1024:
        return Response('request too large', status_code=413)
    return await call_next(request)


HEIF_BRANDS = (b'heic', b'heix', b'hevc', b'hevx', b'heim', b'heis', b'mif1', b'msf1')


def heic_to_jpeg(data):
    # anything else, or an undecodable HEIC, comes back unchanged (then rejected as unsupported)
    if data[4:8] != b'ftyp' or data[8:12] not in HEIF_BRANDS:
        return data
    try:
        img = Image.open(io.BytesIO(data))
        # 48 MP "HEIF Max" photos are refused (about 150 MB per decoded copy; the server has 512 MB).
        # The usual 12/24 MP photos pass. Shrink while decoding if 48 MP receipts turn up.
        if img.width * img.height > 40_000_000:
            print(f'HEIC too large to convert: {img.width}x{img.height}')
            return data
        img = ImageOps.exif_transpose(img).convert('RGB')  # keep the phone's rotation
        buf = io.BytesIO()
        img.save(buf, 'JPEG', quality=90)
        return buf.getvalue()
    except Exception as e:
        print(f'HEIC conversion failed: {e!r}')
        return data


@app.post('/documents', status_code=202)
def upload(files: list[UploadFile], tasks: BackgroundTasks, uid: str = Depends(current_user)):
    # plain def: FastAPI runs it in a thread, so database writes don't block other requests
    if not files or len(files) > MAX_FILES:
        raise HTTPException(400, f'Send 1 to {MAX_FILES} files.')
    with store.conn() as con:
        used = reads_today(con, uid)
    # Counted once per request, so two parallel uploads can go slightly over the limit.
    left = DAILY_UPLOAD_LIMIT - used
    if left <= 0:
        raise HTTPException(429, f"You have used today's {DAILY_UPLOAD_LIMIT} reads. Try again tomorrow.")
    created = []
    for f in files:
        if left <= 0:
            created.append({'file_name': f.filename, 'error': 'Daily limit reached. Try again tomorrow.'})
            continue
        data = f.file.read(MAX_BYTES + 1)
        if len(data) > MAX_BYTES:
            created.append({'file_name': f.filename, 'error': 'Larger than 10 MB.'})
            continue
        data = heic_to_jpeg(data)
        kind = providers.mime(data)  # type from the file's bytes, never from its name
        if kind is None:
            created.append({'file_name': f.filename, 'error': 'Not a PDF, JPG, PNG, WebP or HEIC file.'})
            continue
        if len(data) > MAX_BYTES:  # a HEIC photo that grew past the limit when converted to JPEG
            created.append({'file_name': f.filename, 'error': 'Larger than 10 MB.'})
            continue
        if kind == 'application/pdf':
            pages = pdf_pages(data)
            if not 1 <= pages <= MAX_PDF_PAGES:
                created.append({'file_name': f.filename, 'error': f'PDF must be readable with 1 to {MAX_PDF_PAGES} pages'})
                continue
        with store.conn() as con:
            doc_id = con.execute("INSERT INTO documents (user_id, file_name, mime, file, status) "
                                 "VALUES (%s, %s, %s, %s, 'processing') RETURNING id",
                                 (uid, os.path.basename(f.filename or 'upload'), kind, data)).fetchone()['id']
            con.execute('INSERT INTO reads (user_id) VALUES (%s)', (uid,))
        left -= 1
        tasks.add_task(process, doc_id)
        created.append({'id': doc_id, 'file_name': f.filename, 'status': 'processing'})
    return created


@app.get('/documents')
def list_documents(status: str | None = None, q: str | None = None, date_from: date | None = None,
                   date_to: date | None = None, limit: int = Query(50, ge=1, le=200), offset: int = Query(0, ge=0),
                   oldest: bool = False, uid: str = Depends(current_user)):
    sql = 'SELECT id, file_name, status, vendor, currency, issue_date, total, created_at FROM documents WHERE user_id = %s'
    args = [uid]
    if status:
        sql += ' AND status = %s'; args.append(status)
    if q:
        like = '%' + q.replace('\\', '\\\\').replace('%', r'\%').replace('_', r'\_') + '%'  # % and _ typed are literal
        sql += ' AND (vendor ILIKE %s OR file_name ILIKE %s)'; args += [like] * 2
    if date_from:
        sql += ' AND issue_date >= %s'; args.append(date_from)
    if date_to:
        sql += ' AND issue_date <= %s'; args.append(date_to)
    sql += f' ORDER BY id {"ASC" if oldest else "DESC"} LIMIT %s OFFSET %s'
    with store.conn() as con:
        return con.execute(sql, (*args, limit, offset)).fetchall()


SECOND_FIELDS = ('vendor', 'doc_number', 'issue_date', 'due_date', 'currency', 'subtotal', 'discount', 'tax', 'service_charge', 'total')
LINE_FIELDS = ('quantity', 'unit_price', 'amount', 'discount')


def second_reading_changes(doc: Document, second: Document) -> list[dict]:
    # what the second reading would change, field by field; values the document already has drop off
    text = lambda v: None if v is None else str(v)
    out = []
    for f in SECOND_FIELDS:
        new = getattr(second, f)
        if new is not None and new != getattr(doc, f):
            out.append({'field': f, 'from': text(getattr(doc, f)), 'to': str(new),
                        **({'text': getattr(second, f + '_text')} if f.endswith('_date') else {})})
    if len(second.items) != len(doc.items):
        if second.items:
            out.append({'field': 'items', 'from': f'{len(doc.items)} lines', 'to': [i.model_dump(mode='json') for i in second.items]})
        return out
    for n, (mine, theirs) in enumerate(zip(doc.items, second.items)):
        for f in LINE_FIELDS:
            new = getattr(theirs, f)
            if new is not None and new != getattr(mine, f):
                out.append({'field': f'items[{n}].{f}', 'from': text(getattr(mine, f)), 'to': str(new)})
    return out


MEANINGFUL_TOP = ('subtotal', 'discount', 'tax', 'service_charge', 'total', 'issue_date', 'due_date')


def meaningful_changes(doc: Document, second: Document) -> list[dict]:
    # differences worth a person's look: money, dates, line amounts, and the vendor beyond case and punctuation;
    # 0.00 lines, letter case, currency or branch alone do not count
    def money_lines(d):
        return sorted(v for i in d.items for v in (i.amount, i.discount) if v)

    out = []
    for c in second_reading_changes(doc, second):
        f = c['field']
        if f == 'items':
            if money_lines(doc) != money_lines(second):
                out.append(c)
        elif f in MEANINGFUL_TOP or (f.startswith('items[') and not f.endswith('.quantity')):
            out.append(c)
        elif f == 'vendor' and store.vendor_key(c['from']) != store.vendor_key(c['to']):
            out.append(c)
        elif f == 'doc_number' and c['from'] and store.same_number(c['from']) != store.same_number(c['to']):
            out.append(c)
    return out


@app.get('/documents/{doc_id}')
def get_document(doc_id: int, uid: str = Depends(current_user)):
    with store.conn() as con:
        r = get_row(con, doc_id, uid)
        if r['status'] == 'needs_review' and any(c['check'] == 'duplicate' for c in r['checks'] or []):
            dup = store.duplicate_of(con, uid, Document(**r['document']), before_id=doc_id)
            # the original may have been deleted or corrected since: drop the flag, or link the copy still there
            r['checks'] = [{**c, 'duplicate_of': dup['id']} if c['check'] == 'duplicate' else c
                           for c in r['checks'] if dup or c['check'] != 'duplicate']
        second = r.pop('second_reading')
        has_second = r['status'] == 'needs_review' and second and r['document']
        order = store.date_order(con, uid, r['document']['vendor']) if has_second else None
    r['suggestion'] = suggest(Document(**r['document'])) if r['status'] == 'needs_review' and r['document'] else None
    r['second_reading'] = second_reading_changes(Document(**r['document']), apply_date_order(Document(**second), order)) \
        if has_second else []
    return r


@app.post('/check')
def check(doc: Document, date_order: DateOrderValue | None = None, doc_id: int | None = None,
          uid: str = Depends(current_user)):
    # doc_id: compared only with documents uploaded before it
    doc = doc.model_copy(update={'is_document': None, 'document_count': None})  # the user is editing it: it is one document
    with store.conn() as con:
        order = date_order or store.date_order(con, uid, doc.vendor)
        doc = apply_date_order(doc, order)
        second = doc_id and con.execute('SELECT status, second_reading FROM documents WHERE id = %s AND user_id = %s',
                                        (doc_id, uid)).fetchone()
        changes = second_reading_changes(doc, apply_date_order(Document(**second['second_reading']), order)) \
            if second and second['status'] == 'needs_review' and second['second_reading'] else []
        return {'document': doc, 'checks': run_checks(con, uid, doc, order, doc_id), 'suggestion': suggest(doc),
                'second_reading': changes}


@app.put('/documents/{doc_id}')
def update_document(doc_id: int, doc: Document, date_order: DateOrderValue | None = None,
                    if_unchanged_since: datetime | None = None, uid: str = Depends(current_user)):
    # if_unchanged_since: a change made meanwhile in another tab is not overwritten
    with store.conn() as con:
        row = get_row(con, doc_id, uid)
        if row['status'] == 'processing':
            raise HTTPException(409, 'Still being read. Save again when it is done.')
        if if_unchanged_since and row['updated_at'] != if_unchanged_since:
            raise HTTPException(409, 'Changed in another tab. Reload to see the latest version.')
        doc = doc.model_copy(update={'is_document': None, 'document_count': None})  # the user saved it: it is one document
        save(con, doc_id, doc, 'reviewed', uid, date_order=date_order)
        queue_document_event(con, doc_id)
    WAKE.set()
    return get_document(doc_id, uid)


@app.post('/documents/{doc_id}/retry', status_code=202)
def retry(doc_id: int, tasks: BackgroundTasks, uid: str = Depends(current_user)):
    # reviewed documents keep their corrections; processing ones are not read twice
    with store.conn() as con:  # one statement, so two quick clicks can't both start a read
        if reads_today(con, uid) >= DAILY_UPLOAD_LIMIT:
            raise HTTPException(429, f"You have used today's {DAILY_UPLOAD_LIMIT} reads. Try again tomorrow.")
        claimed = con.execute("UPDATE documents SET status = 'processing', error = NULL, updated_at = now() "
                              "WHERE id = %s AND user_id = %s AND status IN ('failed', 'needs_review') RETURNING id",
                              (doc_id, uid)).fetchone()
        if claimed is None:
            get_row(con, doc_id, uid)  # 404 for ids that are not the user's
            raise HTTPException(409, 'Only failed or flagged documents can be read again.')
        con.execute('INSERT INTO reads (user_id) VALUES (%s)', (uid,))
    tasks.add_task(process, doc_id)
    return {'id': doc_id, 'status': 'processing'}


@app.delete('/documents/{doc_id}', status_code=204)
def delete_document(doc_id: int, uid: str = Depends(current_user)):
    with store.conn() as con:
        gone = con.execute('DELETE FROM documents WHERE id = %s AND user_id = %s RETURNING status', (doc_id, uid)).fetchone()
        if gone and gone['status'] in ('passed', 'reviewed'):  # only checked documents have a spreadsheet row to mark
            queue_event(con, uid, {'event': 'document.deleted', 'id': doc_id, 'status': 'deleted'})
    if gone is None:
        raise HTTPException(404, 'Document not found.')
    WAKE.set()


@app.get('/documents/{doc_id}/file')
def get_file(doc_id: int, uid: str = Depends(current_user)):
    with store.conn() as con:
        r = get_row(con, doc_id, uid, with_file=True)
    # RFC 5987 form: any language in the name, and no quotes or line breaks can reach the header
    return Response(bytes(r['file']), media_type=r['mime'],
                    headers={'Content-Disposition': f"inline; filename*=UTF-8''{urllib.parse.quote(r['file_name'])}"})


def pdf_pages(data):
    with PDF_LOCK:
        try:
            pdf = pypdfium2.PdfDocument(bytes(data))
        except pypdfium2.PdfiumError:
            return 0
        try:
            return len(pdf)
        finally:
            pdf.close()


@app.get('/documents/{doc_id}/pages')
def page_count(doc_id: int, uid: str = Depends(current_user)):
    with store.conn() as con:
        r = get_row(con, doc_id, uid, with_file=True)
    return {'pages': pdf_pages(r['file']) if r['mime'] == 'application/pdf' else 1}


@app.get('/documents/{doc_id}/pages/{n}')
def page_image(doc_id: int, n: int, uid: str = Depends(current_user)):
    with store.conn() as con:
        r = get_row(con, doc_id, uid, with_file=True)
    if r['mime'] != 'application/pdf':
        if n != 0:
            raise HTTPException(404, 'Page not found.')
        return Response(bytes(r['file']), media_type=r['mime'])
    buf = io.BytesIO()
    with PDF_LOCK:
        pdf = pypdfium2.PdfDocument(bytes(r['file']))
        try:
            if not 0 <= n < len(pdf):
                raise HTTPException(404, 'Page not found.')
            page = pdf[n]
            scale = min(2, 2000 / max(page.get_size()))  # at most 2000 px on the long side, whatever the page size
            page.render(scale=scale).to_pil().save(buf, 'PNG')
        finally:
            pdf.close()
    return Response(buf.getvalue(), media_type='image/png')


class DateOrder(BaseModel):
    date_order: DateOrderValue


@app.put('/vendors/{vendor:path}/date-order')  # :path keeps names like 'M/S Rahman Traders'
def set_vendor_date_order(vendor: str, body: DateOrder, uid: str = Depends(current_user)):
    with store.conn() as con:
        store.set_date_order(con, uid, vendor, body.date_order)
        rows = [r for r in con.execute("SELECT id, vendor, document FROM documents WHERE user_id = %s "
                                       "AND status = 'needs_review'", (uid,)).fetchall()
                if store.vendor_key(r['vendor']) == store.vendor_key(vendor)]
        for r in rows:
            save(con, r['id'], Document.model_validate(r['document']), None, uid)
            if row_by_id(con, r['id'])['status'] == 'passed':
                queue_document_event(con, r['id'])
    WAKE.set()
    return {'vendor': vendor, 'date_order': body.date_order, 'rechecked': len(rows)}


def delete_auth_user(uid):
    key = os.environ.get('SUPABASE_SECRET_KEY')
    if not key:
        raise HTTPException(503, 'Account deletion is not set up on this server.')
    url = f'{os.environ["SUPABASE_URL"].rstrip("/")}/auth/v1/admin/users/{uid}'
    try:
        urllib.request.urlopen(urllib.request.Request(url, method='DELETE', headers={'apikey': key, 'Authorization': f'Bearer {key}'}),
                               timeout=15)
    except urllib.error.HTTPError as e:
        if e.code != 404:  # 404: already removed by an earlier try whose answer was lost, so a retry still finishes
            raise


@app.delete('/account', status_code=204)
def delete_account(uid: str = Depends(current_user)):
    # if the sign-in account can't be removed, the database changes roll back
    with store.conn() as con:
        con.execute('DELETE FROM documents WHERE user_id = %s', (uid,))
        con.execute('DELETE FROM vendor_date_orders WHERE user_id = %s', (uid,))
        con.execute('DELETE FROM reads WHERE user_id = %s', (uid,))
        con.execute('DELETE FROM webhooks WHERE user_id = %s', (uid,))
        con.execute('DELETE FROM webhook_events WHERE user_id = %s', (uid,))
        try:
            delete_auth_user(uid)
        except HTTPException:
            raise
        except Exception as e:
            print(f'deleting auth user {uid} failed: {e!r}')
            raise HTTPException(502, 'Could not delete the account. Please try again.')


class WebhookIn(BaseModel):
    url: str = Field(max_length=2000)


@app.get('/account/webhook')
def get_webhook(uid: str = Depends(current_user)):
    # the secret is never sent back: the user sees it once, when it is made
    with store.conn() as con:
        target = webhook_for(con, uid)
        r = con.execute('SELECT last_at, last_error FROM webhooks WHERE user_id = %s', (uid,)).fetchone()
    own = bool(target and target[2])
    error = r and r['last_error']
    if r and not own and fernet():  # WEBHOOK_KEY changed: the saved address can't be read any more
        error = 'Your saved address could not be read after a server change. Save it again.'
    return {'enabled': bool(fernet()), 'url': target[0] if own else None, 'last_at': r and r['last_at'], 'last_error': error}


@app.put('/account/webhook')
def set_webhook(body: WebhookIn, uid: str = Depends(current_user)):
    if not fernet():
        raise HTTPException(503, 'Webhooks are not set up on this server.')
    url = body.url.strip()
    if not public_url(url):
        raise HTTPException(400, 'Use an https address that is reachable on the internet.')
    secret = secrets.token_urlsafe(32)  # new on every save; to replace a leaked one, remove the address and save it again
    with store.conn() as con:
        con.execute('INSERT INTO webhooks (user_id, url, secret) VALUES (%s, %s, %s) ON CONFLICT (user_id) DO UPDATE '
                    'SET url = excluded.url, secret = excluded.secret, last_at = NULL, last_error = NULL',
                    (uid, fernet().encrypt(url.encode()).decode(), fernet().encrypt(secret.encode()).decode()))
    return {'url': url, 'secret': secret}


@app.delete('/account/webhook', status_code=204)
def delete_webhook(uid: str = Depends(current_user)):
    with store.conn() as con:
        if con.execute('DELETE FROM webhooks WHERE user_id = %s', (uid,)).rowcount:  # the server's n8n accounts keep theirs
            con.execute('DELETE FROM webhook_events WHERE user_id = %s', (uid,))


TESTED = {}  # user -> time of the last Send test (in memory: a restart allows one more)


@app.post('/account/webhook/test')
def test_webhook(uid: str = Depends(current_user)):
    with store.conn() as con:
        target = webhook_for(con, uid)
    if not (target and target[2]):
        raise HTTPException(404, 'Add a webhook address first.')
    if time.monotonic() - TESTED.get(uid, -60) < 10:  # each test is a request this server makes for the user
        raise HTTPException(429, 'Wait a few seconds before testing again.')
    TESTED[uid] = time.monotonic()
    try:
        return {'ok': True, 'status': open_webhook(target[0], {'event': 'test', 'id': None}, target[1], True, 30)}
    except Exception as e:
        return {'ok': False, 'error': str(e)[:200]}


@app.get('/usage')
def usage(uid: str = Depends(current_user)):
    with store.conn() as con:
        return {'used': reads_today(con, uid), 'limit': DAILY_UPLOAD_LIMIT}


@app.get('/stats')
def stats(date_from: date | None = None, uid: str = Depends(current_user)):
    # money per currency (never converted), only from checked documents
    with store.conn() as con:
        q = lambda cols, rest='', args=(): con.execute(f'SELECT {cols} FROM documents WHERE user_id = %s {rest}', (uid, *args)).fetchall()
        checked = "AND status IN ('passed', 'reviewed')" + (' AND issue_date >= %s' if date_from else '')
        m = lambda cols, rest: q(cols, f'{checked} {rest}', (date_from,) if date_from else ())
        return {
            'documents': q('count(*) AS n')[0]['n'],
            'by_status': {r['status']: r['n'] for r in q('status, count(*) AS n', 'GROUP BY status')},
            'spend_by_currency': m('currency, sum(total) AS total, count(*) AS n', 'AND total IS NOT NULL GROUP BY currency'),
            'tax_by_currency': m("currency, sum((document->>'tax')::numeric) AS total, count(*) AS n",
                                 "AND document->>'tax' IS NOT NULL GROUP BY currency"),
            'top_vendors': m('min(vendor) AS vendor, currency, sum(total) AS total, count(*) AS n',
                             'AND vendor IS NOT NULL GROUP BY lower(vendor), currency ORDER BY total DESC NULLS LAST'),
            'by_month': m("to_char(issue_date, 'YYYY-MM') AS month, currency, sum(total) AS total, count(*) AS n",
                          'AND issue_date IS NOT NULL GROUP BY month, currency ORDER BY month'),
        }


DOC_COLUMNS = ['id', 'file_name', 'status', 'doc_type', 'vendor', 'branch', 'buyer', 'doc_number', 'issue_date', 'due_date',
               'currency', 'subtotal', 'discount', 'tax', 'service_charge', 'total']
ITEM_COLUMNS = ['document_id', 'description', 'quantity', 'unit_price', 'amount', 'discount']


def export_rows(uid, status):
    sql = "SELECT id, file_name, status, document FROM documents WHERE user_id = %s AND status != 'failed' AND document IS NOT NULL"
    args = [uid]
    if status:
        sql += ' AND status = %s'; args.append(status)
    docs, items = [], []
    with store.conn() as con:
        for r in con.execute(sql + ' ORDER BY id', args):
            d = r['document']
            docs.append([r['id'], r['file_name'], r['status']] + [d.get(k) for k in DOC_COLUMNS[3:]])
            items += [[r['id']] + [i.get(k) for k in ITEM_COLUMNS[1:]] for i in d.get('items') or []]
    return docs, items


NUMERIC = {'subtotal', 'discount', 'tax', 'service_charge', 'total', 'quantity', 'unit_price', 'amount'}


def cell(column, v, xlsx=False):
    # numbers stay numbers, text stays text (00123 keeps its zeros);
    # text starting with = + - @ tab or CR gets a leading ' so a spreadsheet never runs it
    if v is None:
        return None
    if column in NUMERIC:
        return Decimal(format(v if isinstance(v, Decimal) else Decimal(str(v)), 'f'))  # 10, not 1E+1
    if isinstance(v, str) and v[:1] in (('=',) if xlsx else ('=', '+', '-', '@', '\t', '\r')):
        return "'" + v
    return v


def cells(columns, row_, xlsx=False):
    return [cell(c, v, xlsx) for c, v in zip(columns, row_)]


# QuickBooks Online "Import bills" layout; headers are matched to QuickBooks fields during the import.
QB_COLUMNS = ['Bill no.', 'Supplier', 'Bill Date', 'Due Date', 'Account', 'Line Description', 'Line Amount', 'Line Tax Code']


def quickbooks_rows(uid):
    # bill lines add up to the total, else one line with the total;
    # documents without a date or total are skipped (QuickBooks needs both)
    rows, skipped = [], 0
    with store.conn() as con:
        found = con.execute("SELECT id, document FROM documents WHERE user_id = %s AND status IN ('passed', 'reviewed') "
                            'ORDER BY id', (uid,))
        for r in found:
            d = Document.model_validate(r['document'])
            if d.issue_date is None or d.total is None:
                skipped += 1
                continue
            items = [(i.description, i.amount - abs(i.discount or 0)) for i in d.items if i.amount is not None]
            extra = [(name, v) for name, v in (('Service charge', d.service_charge), ('Discount', -abs(d.discount) if d.discount else None)) if v]
            # tax added on top gets its own line; "VAT included" is already inside the items. Whichever adds up wins.
            lines = next((ls for ls in (items + extra + ([('Tax', d.tax)] if d.tax else []), items + extra)
                          if items and sum(v for _, v in ls) == d.total), [('Total', d.total)])
            # US date order; QuickBooks asks for the file's date format on import.
            bill = [cell('text', d.doc_number or f'CC-{r["id"]}'), cell('text', d.vendor or 'Unknown supplier'),
                    f'{d.issue_date:%m/%d/%Y}', f'{(d.due_date or d.issue_date):%m/%d/%Y}', 'Uncategorized Expense']
            # 2 decimals, or 3 for currencies such as KWD
            rows += [bill + [cell('text', desc), f'{v:.2f}' if v == round(v, 2) else format(v.normalize(), 'f'), None] for desc, v in lines]
    return rows, skipped


@app.get('/export')
def export(format: str = 'xlsx', status: str | None = None, uid: str = Depends(current_user)):
    # X-Skipped header: documents left out of the QuickBooks file
    if format not in ('csv', 'xlsx', 'quickbooks'):
        raise HTTPException(422, 'Format must be csv, xlsx or quickbooks.')
    if format == 'quickbooks':
        rows, skipped = quickbooks_rows(uid)
        buf = io.StringIO()
        csv.writer(buf).writerows([QB_COLUMNS] + rows)
        return Response(buf.getvalue().encode('utf-8-sig'), media_type='text/csv', headers={
            'Content-Disposition': 'attachment; filename="quickbooks-bills.csv"', 'X-Skipped': str(skipped)})
    docs, items = export_rows(uid, status)
    if format == 'csv':
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(DOC_COLUMNS)
        w.writerows(cells(DOC_COLUMNS, r) for r in docs)
        return Response(buf.getvalue().encode('utf-8-sig'), media_type='text/csv',
                        headers={'Content-Disposition': 'attachment; filename="documents.csv"'})
    from openpyxl import Workbook
    wb = Workbook()
    for ws, cols, rows in ((wb.active, DOC_COLUMNS, docs), (wb.create_sheet('Items'), ITEM_COLUMNS, items)):
        ws.append(cols)
        for row_ in rows:
            ws.append([date.fromisoformat(v) if c in ('issue_date', 'due_date') and v else v for c, v in zip(cols, cells(cols, row_, xlsx=True))])
        for row_ in ws.iter_rows(min_row=2):  # real dates, so Excel's date filters and sorting work
            for x in row_:
                if isinstance(x.value, date):
                    x.number_format = 'yyyy-mm-dd'
    wb.active.title = 'Documents'
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition': 'attachment; filename="documents.xlsx"'})
